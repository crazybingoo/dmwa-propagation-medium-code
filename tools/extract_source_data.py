"""Export public workbook blocks as CSV (requires openpyxl).

These are panel source blocks, not the richer intermediates used by all legacy scripts.
"""
from pathlib import Path
import argparse,csv,re
from openpyxl import load_workbook

def main():
    root=Path(__file__).resolve().parents[1]
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,default=root/'Source_Data.xlsx')
    parser.add_argument('--output',type=Path,default=root/'outputs/source_blocks')
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    wb=load_workbook(args.workbook,read_only=True,data_only=True)
    for index,row in enumerate(wb['Data_Map'].iter_rows(min_row=2,values_only=True),1):
        sheet,figure,block,description,unit,source,cell_range,n=row[:8]
        if not sheet:continue
        stem=re.sub(r'[^A-Za-z0-9_-]+','_',f'{sheet}_{block}')
        path=args.output/f'{index:02d}_{stem}.csv'
        with path.open('w',newline='',encoding='utf-8') as f:
            csv.writer(f).writerows([[c.value for c in r] for r in wb[sheet][cell_range]])
        print(path.name)
    wb.close()

if __name__=='__main__':main()
